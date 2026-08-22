from __future__ import annotations

from io import BytesIO
from pathlib import Path
import textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from build_scientific_article import (
    ACCENT,
    ACCENT_DARK,
    GRID,
    INK,
    MUTED,
    PALE_ALT,
    WHITE,
    add_hyperlink,
    add_page_number,
    format_cell,
    paragraph_border_bottom,
    set_cell_shading,
    set_repeat_table_header,
    set_run_font,
    set_table_borders,
    set_table_geometry,
)


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "pressq_quiu_conference_paper_en_accessible.docx"


def conference_architecture_diagram() -> BytesIO:
    width, height = 1900, 940
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)

    def font(size: int, bold: bool = False):
        candidates = [
            Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
            Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
        ]
        for candidate in candidates:
            if candidate.exists():
                return ImageFont.truetype(str(candidate), size=size)
        return ImageFont.load_default()

    title_font = font(31, True)
    body_font = font(24)
    note_font = font(20)
    dark = "#29345F"

    def box(coords, fill, title, body, wrap=25):
        draw.rounded_rectangle(coords, radius=20, fill=fill, outline=dark, width=3)
        x1, y1, x2, _ = coords
        draw.text(((x1 + x2) / 2, y1 + 30), title, font=title_font, fill=dark, anchor="ma")
        wrapped = "\n".join(textwrap.fill(line, width=wrap) for line in body.splitlines())
        draw.multiline_text(
            ((x1 + x2) / 2, y1 + 88),
            wrapped,
            font=body_font,
            fill="#343A48",
            anchor="ma",
            align="center",
            spacing=7,
        )

    def arrow(x1, y1, x2, y2):
        draw.line((x1, y1, x2 - 25, y2), fill=dark, width=5)
        draw.polygon([(x2 - 28, y2 - 14), (x2, y2), (x2 - 28, y2 + 14)], fill=dark)

    box((45, 110, 390, 335), "#E8EDF9", "SOURCE LAYER", "LGBTQ Video Game Archive\nsupplementary sources\ncurrent curated entries", 23)
    box((455, 110, 825, 335), "#F0E9F7", "STRUCTURED DATA", "characters\ngame systems\nqueer readings\nsource details", 22)
    box((890, 110, 1260, 335), "#E9F4F0", "PLANNED CALCULATOR", "checked filters\nmatching / total records\nmulti-axis results", 20)
    box((1325, 110, 1685, 335), "#EEF1F8", "QUIU", "data-grounded explanation\nqualifications\nsource-aware answers", 23)
    box((1325, 470, 1685, 675), "#FFF3E7", "RESEARCHER", "asks questions\nchecks evidence\ninterprets dataset results", 23)

    arrow(390, 222, 455, 222)
    arrow(825, 222, 890, 222)
    arrow(1260, 222, 1325, 222)
    arrow(1505, 335, 1505, 470)

    draw.text(
        (950, 790),
        "The dataset guides Quiu's answer; it does not train the language model.",
        font=note_font,
        fill="#5F6368",
        anchor="mm",
    )
    draw.text(
        (950, 835),
        "For publishable statistics, a separate tool should check the calculation before Quiu explains it.",
        font=note_font,
        fill="#5F6368",
        anchor="mm",
    )

    stream = BytesIO()
    image.save(stream, format="PNG", dpi=(220, 220))
    stream.seek(0)
    return stream


def add_alt_text(shape, title: str, description: str):
    doc_pr = shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    section.different_first_page_header_footer = False

    settings = doc.settings._element
    even_odd = settings.find(qn("w:evenAndOddHeaders"))
    if even_odd is None:
        even_odd = OxmlElement("w:evenAndOddHeaders")
        settings.append(even_odd)
    even_odd.set(qn("w:val"), "true")

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.4)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(0.7)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.14

    for name, size, before, after in (
        ("Heading 1", 12.8, 12, 5),
        ("Heading 2", 11.4, 9, 4),
        ("Heading 3", 10.5, 7, 3),
    ):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(ACCENT_DARK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.line_spacing = 1.0

    abstract_style = styles.add_style("Conference Abstract", WD_STYLE_TYPE.PARAGRAPH)
    abstract_style.font.name = "Times New Roman"
    abstract_style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    abstract_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    abstract_style.font.size = Pt(9.6)
    abstract_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abstract_style.paragraph_format.first_line_indent = Cm(0)
    abstract_style.paragraph_format.space_after = Pt(4)
    abstract_style.paragraph_format.line_spacing = 1.03

    caption_style = styles.add_style("Conference Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption_style.font.name = "Times New Roman"
    caption_style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    caption_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    caption_style.font.size = Pt(9.1)
    caption_style.font.bold = True
    caption_style.font.color.rgb = RGBColor.from_string(ACCENT_DARK)
    caption_style.paragraph_format.first_line_indent = Cm(0)
    caption_style.paragraph_format.space_before = Pt(6)
    caption_style.paragraph_format.space_after = Pt(3)
    caption_style.paragraph_format.keep_with_next = True

    source_style = styles.add_style("Conference Source", WD_STYLE_TYPE.PARAGRAPH)
    source_style.font.name = "Times New Roman"
    source_style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    source_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    source_style.font.size = Pt(8.2)
    source_style.font.color.rgb = RGBColor.from_string(MUTED)
    source_style.paragraph_format.first_line_indent = Cm(0)
    source_style.paragraph_format.space_before = Pt(3)
    source_style.paragraph_format.space_after = Pt(5)
    source_style.paragraph_format.line_spacing = 1.0

    for header in (section.header, section.even_page_header):
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("PRESS Q + QUIU | CONFERENCE MANUSCRIPT")
        set_run_font(r, size=8.2, bold=True, color=MUTED)

    for footer in (section.footer, section.even_page_footer):
        add_page_number(footer.paragraphs[0])

    doc.core_properties.title = "Press Q and Quiu: A Data-Grounded Research Tool for LGBTQ+ Video Game Studies"
    doc.core_properties.subject = "Conference manuscript on structured data, clear statistics, and evidence-based AI assistance"
    doc.core_properties.keywords = "LGBTQ+; video games; research infrastructure; grounded AI; digital humanities; queer archives"

    def add_heading(text: str, level=1):
        return doc.add_paragraph(text, style=f"Heading {level}")

    def add_body(text: str, indent=True, keep=False):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.7) if indent else Cm(0)
        p.paragraph_format.keep_with_next = keep
        r = p.add_run(text)
        set_run_font(r, size=10.4, color=INK)
        return p

    def add_abstract(label: str, text: str):
        p = doc.add_paragraph(style="Conference Abstract")
        lead = p.add_run(label + ": ")
        set_run_font(lead, size=9.6, bold=True, color=ACCENT_DARK)
        body = p.add_run(text)
        set_run_font(body, size=9.6, color=INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run("PRESS Q AND QUIU: BUILDING A RESEARCH TOOL FOR LGBTQ+ VIDEO GAME STUDIES")
    set_run_font(r, size=15.5, bold=True, color=ACCENT_DARK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(13)
    r = p.add_run("A work-in-progress tool for clear statistics, intersectional analysis, and evidence-based answers")
    set_run_font(r, size=11.1, italic=True, color=MUTED)

    for text, bold in (
        ("[AUTHOR INFORMATION REMOVED FOR CONFERENCE REVIEW]", True),
        ("[Affiliation | ORCID | Corresponding author e-mail]", False),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(2 if bold else 10)
        r = p.add_run(text)
        set_run_font(r, size=10.1, bold=bold, color=INK if bold else MUTED)

    rule = doc.add_paragraph()
    rule.paragraph_format.first_line_indent = Cm(0)
    rule.paragraph_format.space_after = Pt(7)
    paragraph_border_bottom(rule, color="AEB7D2", size="8", space="4")

    add_abstract(
        "Abstract",
        "This work-in-progress paper presents Press Q and its conversational assistant, Quiu. The project is being built to help researchers study LGBTQ+ representation in video games. Press Q organizes evidence into three types of records: characters, game systems, and queer readings. The long-term goal is to let researchers filter the data, create visualizations, calculate clearly defined statistics, study intersectionality, and ask questions in everyday language (PRESS Q CODE, 2026; PRESS Q DATASET, 2026). We reviewed the source code and a provisional dataset of 94 records: 60 character records, 22 game-system records, and 12 queer-reading records (PRESS Q CODE, 2026; PRESS Q DATASET, 2026). Our test calculations show that some questions can be answered within this dataset when it is clear what is being counted and which records are included. These calculations do not show that automatic statistics are already complete, and they cannot describe the game industry as a whole (PRESS Q DATASET, 2026; LGBTQ VIDEO GAME ARCHIVE, 2025a). Intersectional analysis is a major goal because the Archive’s public website gives limited attention to areas such as race, class, and disability. However, the current Press Q fields are not yet consistent enough for reliable intersectional percentages (LGBTQ VIDEO GAME ARCHIVE, 2025a; PRESS Q DATASET, 2026). Quiu uses the dataset as context; the dataset does not train the model. This approach may reduce unsupported answers, but it cannot guarantee accuracy (PRESS Q CODE, 2026; LEWIS et al., 2020; SHUSTER et al., 2021; NIST, 2024).",
    )
    add_abstract("Keywords", "LGBTQ+ video game representation; intersectionality; research infrastructure; evidence-based AI support; digital humanities; queer archives; research data")

    add_heading("1 Introduction", 1)
    add_body("The LGBTQ Video Game Archive makes decades of LGBTQ+ game content easier to find. It brings together evidence from games, creator statements, journalism, walkthroughs, forums, wikis, videos, criticism, and fan interpretations. The Archive describes itself as a curated collection of information and a starting point for further research, rather than as a traditional archive of primary sources (LGBTQ VIDEO GAME ARCHIVE, 2025a).")
    add_body("Finding examples is not the same as calculating statistics. The Archive uses WordPress categories that can overlap. A post may also discuss several examples or appear in more than one category. This is useful for browsing, but the visible category totals do not always represent the same type of item (LGBTQ VIDEO GAME ARCHIVE, 2025a, 2025b). The Archive lists more than 1,200 games, but only about 400 had been fully researched by June 2025. It also states that there is no reliable count of all games released each year. For this reason, it cannot calculate the percentage of all games that contain LGBTQ+ content (LGBTQ VIDEO GAME ARCHIVE, 2025a).")
    add_body("Press Q is being developed for a related research need. It turns selected evidence and new entries into structured records that researchers can inspect. The project may later support filters, tables, visualizations, counts, percentages within the dataset, and intersectional comparisons. Some parts already work, but automatic calculations, permanent data storage, consistent categories, and user testing are still under development. Quiu is an early question-and-answer interface. Its instructions tell it to use the Press Q data, show when information is missing, identify sources, keep disputed evidence visible, and avoid mixing characters, game systems, and queer readings (PRESS Q CODE, 2026; PRESS Q DATASET, 2026).")
    add_body("This paper has four aims. First, it explains how Press Q organizes its data and how Quiu is intended to support research. Second, it tests which percentage questions the current data could support without claiming that the feature is finished. Third, it explains what the project still needs before it can calculate intersectional statistics. Fourth, it compares Press Q with gaps identified in the LGBTQ Video Game Archive while treating the two projects as complementary (PRESS Q CODE, 2026; PRESS Q DATASET, 2026; LGBTQ VIDEO GAME ARCHIVE, 2025a).")

    add_heading("1.1 Research questions", 2)
    questions = [
        "RQ1. Which counts and percentages could the developing Press Q dataset support in the future?",
        "RQ2. How can Quiu use the dataset to support research and reduce unsupported answers?",
        "RQ3. Which gaps in the LGBTQ Video Game Archive could Press Q help address, and which gaps would remain?",
        "RQ4. What information and quality checks are needed before Press Q can calculate intersectional statistics?",
    ]
    for question in questions:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(question)
        set_run_font(r, size=10.4, bold=True, color=INK)

    add_heading("2 Background and Related Work", 1)
    add_heading("2.1 Queer game archives and what gets counted", 2)
    add_body("Queer game research studies more than characters whose identities are directly stated. Researchers also examine play, game rules, authorship, audience response, and interpretation. A queer reading can be historically important even when it does not prove that a character has a canonical LGBTQ+ identity (SHAW, 2014; RUBERG; SHAW, 2017; RUBERG, 2019; SHAW; PERSAUD, 2020). The Archive reflects this broad approach. It records explicit and implied representation, player-created modifications, background content, homophobia and transphobia, and material that players or critics have read as queer (LGBTQ VIDEO GAME ARCHIVE, 2025a, 2025b).")
    add_body("This broad approach creates a counting problem: researchers must decide what kind of item they are counting. For example, a romance system shows what a player can choose, but it does not prove a fixed identity for every character involved. A queer reading records how a game has been interpreted, but it does not automatically become a fact about the story. Previous research shows that statistics and histories of interpretation can be studied together when these differences remain clear (SHAW et al., 2019; SHAW; PERSAUD, 2020).")

    add_heading("2.2 Data choices, sources, and intersectionality", 2)
    add_body("Data in the humanities are shaped by decisions about what to include and how to classify it. A chart or percentage can hide these decisions if it does not show what was counted, what is missing, and what remains uncertain (DRUCKER, 2011; D'IGNAZIO; KLEIN, 2020). The FAIR principles also stress the importance of clear descriptions, sources, and reuse information when research data are shared (WILKINSON et al., 2016).")
    add_body("Intersectionality examines how identities and structures of power work together. Looking at only one category at a time can hide experiences created by the overlap of gender, sexuality, race, class, disability, and other factors (CRENSHAW, 1989). The Archive’s public website mainly codes gender and sexuality. It also acknowledges that race, class, disability, character role, and cultural differences need more research. Researchers can request its underlying spreadsheet when they need more detail (LGBTQ VIDEO GAME ARCHIVE, 2025a). Press Q treats this as an opportunity for future development. The aim is to make combinations of well-sourced identity information searchable without guessing missing identities or treating “not documented” as “not present” (D'IGNAZIO; KLEIN, 2020; PRESS Q DATASET, 2026).")

    add_heading("2.3 Giving an AI evidence is not the same as training it", 2)
    add_body("Language models can use outside sources when they answer a question. Research shows that this can improve factual accuracy in some tasks, especially when a system retrieves relevant evidence before generating an answer (LEWIS et al., 2020). Similar methods have reduced unsupported statements in some dialogue studies. However, these results do not guarantee accuracy for every model, dataset, question, or research field (SHUSTER et al., 2021).")
    add_body("Quiu uses a simpler approach. For each question, the application places all three CSV files in the model’s instructions. It does not search an index for selected records, and it does not change or fine-tune the model itself (PRESS Q CODE, 2026). NIST uses the term confabulation for incorrect content that a generative AI system presents with confidence. Adding context may help, but it does not remove this risk. The system must still be tested and monitored (NIST, 2024).")

    add_heading("3 Press Q and Quiu", 1)
    add_heading("3.1 Types of records", 2)
    add_body("Press Q separates the data into three record types. Character records describe a character and the evidence about that character’s identity or coding. Game-system records describe choices offered by a game, such as gender customization, same-gender marriage, romance that does not depend on gender, or queer family creation. Queer-reading records document interpretations, debates, audience responses, and evidence that may challenge a reading. The current dataset has 28 fields for character records and 17 fields for each of the other two record types (PRESS Q DATASET, 2026).")
    add_body("The files record the source, source language, how the item was found, research status, confidence in the evidence, platform or game version, last review date, notes, and other details. Character records also have early fields for whether intersectionality is present and how it is described. These fields help researchers trace where information came from and how it changed over time. They also keep confirmed identity, player choice, and audience interpretation separate. However, the intersectionality fields are not yet consistent enough for statistical analysis (PRESS Q DATASET, 2026; WILKINSON et al., 2016).")

    add_heading("3.2 Adding and reviewing records", 2)
    add_body("The current prototype combines starter records based on archival research with entries reviewed through the Press Q administration page. An administrator can create, edit, delete, and export character and game-system records. At present, these changes are written to local CSV files (PRESS Q CODE, 2026). The public contribution page creates a text draft that a contributor can copy or download. It does not place an unreviewed public submission directly into the research data (PRESS Q CODE, 2026).")
    add_body("This review step protects the data from automatic public changes, but the system is not yet a production database. The administration page warns that files may not remain stored on some serverless hosting services and lists a database connection as a future step (PRESS Q CODE, 2026). In this paper, “database” refers to the structured research data and the permanent storage system that the project plans to build. The version reviewed for this study still uses CSV files, and several research features remain planned rather than complete (PRESS Q CODE, 2026).")

    add_heading("3.3 Current and planned architecture", 2)
    add_body("The current and planned systems are different. At present, the chat service gives all records to the language model, which may be asked to count or summarize them. Figure 1 shows the proposed next step. A separate calculation tool would first check the filters and data values, then return the number of matching records and the total number examined. It would calculate intersectional results only when enough reliable data are available. Quiu would then explain the checked result and point to the relevant sources (PRESS Q CODE, 2026; NIST, 2024).")
    cap = doc.add_paragraph("Figure 1 - Proposed research-support architecture for Press Q and Quiu", style="Conference Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p = doc.add_paragraph()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.paragraph_format.first_line_indent = Cm(0)
    pic_p.paragraph_format.space_after = Pt(2)
    shape = pic_p.add_run().add_picture(conference_architecture_diagram(), width=Inches(5.8))
    add_alt_text(
        shape,
        "Press Q and Quiu architecture",
        "Flow diagram from archival and current sources to structured data, a planned calculation tool, Quiu, and the researcher. A note states that the dataset guides the answer rather than training the model.",
    )
    doc.add_paragraph("Source: authors’ design based on the Press Q code and dataset snapshots (PRESS Q CODE, 2026; PRESS Q DATASET, 2026).", style="Conference Source")

    add_heading("3.4 Current safeguards in Quiu", 2)
    add_body("Quiu’s instructions set several limits. It should use only the Press Q data supplied with the question, reply in the user’s language, and keep characters, game systems, and queer readings separate. It should preserve disputed readings, avoid guessing identity from game rules, explain which records were included, distinguish missing information from a documented absence, and refuse claims about the whole game industry (PRESS Q CODE, 2026). The reviewed version uses a temperature setting of 0.35. This may make the wording more stable, but it does not prove that the answers are accurate (PRESS Q CODE, 2026; NIST, 2024).")

    add_heading("4 Method", 1)
    add_heading("4.1 Case study", 2)
    add_body("This exploratory case study reviewed the Press Q project as it existed on 22 August 2026. We examined the Next.js website, the Quiu API route, the public contribution page, the administration page, the project methodology, and three CSV files: pressq_seed_dataset.csv, game_queer_systems.csv, and queer_readings.csv (PRESS Q CODE, 2026; PRESS Q DATASET, 2026).")
    add_body("We also reviewed the LGBTQ Video Game Archive’s public methodology, stated limitations, category descriptions, and related publications. The study did not review every Archive entry or examine the Archive team’s internal spreadsheet (LGBTQ VIDEO GAME ARCHIVE, 2025a, 2025b; SHAW et al., 2019; SHAW; PERSAUD, 2020).")

    add_heading("4.2 Testing which questions the data can answer", 2)
    add_body("We calculated several example percentages outside the application. These were tests of the dataset, not results from a finished automatic feature. We treated a question as provisionally answerable only when the data clearly defined the matching records, the total group, and the values being compared. Each percentage uses n/N × 100 and describes only the Press Q snapshot reviewed for this paper (PRESS Q DATASET, 2026).")
    add_body("We tested five questions: confirmation status, playable status, romance that does not depend on gender, representation across the whole game industry, and intersectional data coverage. The first three have totals within the local dataset and show what a future calculation tool could produce. The industry question has no reliable total for comparison. The intersectionality question cannot yet be answered because the current fields use inconsistent values and do not record multiple identity dimensions in a consistent way (PRESS Q DATASET, 2026; LGBTQ VIDEO GAME ARCHIVE, 2025a).")

    add_heading("4.3 Limits of the AI assessment", 2)
    add_body("We reviewed Quiu’s code and instructions, but we did not run a formal model evaluation. We did not measure unsupported answers, calculation errors, performance across languages, citation accuracy, or responses to misleading questions. The study therefore examines the safeguards in the design; it does not claim that Quiu has reached a tested level of accuracy (PRESS Q CODE, 2026; NIST, 2024).")

    add_heading("5 Preliminary Findings", 1)
    add_heading("5.1 Dataset composition", 2)
    add_body("The snapshot contains 94 records linked to 38 different game or series names. The same title can appear in more than one file, so 38 is the number of unique names across all three files, not the sum of the title counts in each file (PRESS Q DATASET, 2026).")
    doc.add_paragraph("Table 1 - Composition of the Press Q snapshot", style="Conference Caption")
    table1 = doc.add_table(rows=1, cols=5)
    headers1 = ["Record type", "Records", "Unique titles", "Years", "What the record describes"]
    for idx, text in enumerate(headers1):
        format_cell(table1.rows[0].cells[idx], text, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_run_font(table1.rows[0].cells[idx].paragraphs[0].runs[0], size=8.6, bold=True, color=WHITE)
    set_repeat_table_header(table1.rows[0])
    rows1 = [
        ("Characters", "60", "24", "1990–2020", "Evidence about a character’s identity or coding"),
        ("Game systems", "22", "14", "1990–2018", "Choices and possibilities offered to players"),
        ("Queer readings", "12", "7", "1984–2016", "Interpretations, responses, and opposing evidence"),
    ]
    for row_data in rows1:
        row = table1.add_row()
        for idx, text in enumerate(row_data):
            format_cell(row.cells[idx], text, align=WD_ALIGN_PARAGRAPH.LEFT if idx in (0, 4) else WD_ALIGN_PARAGRAPH.CENTER)
            set_run_font(row.cells[idx].paragraphs[0].runs[0], size=8.5, color=INK)
    set_table_geometry(table1, [1900, 1100, 1250, 1250, 3572])
    set_table_borders(table1)
    doc.add_paragraph("Source: Press Q dataset snapshot, 22 August 2026 (PRESS Q DATASET, 2026).", style="Conference Source")

    add_heading("5.2 Which calculations could the dataset support?", 2)
    add_body("In the 60 character records, 41 are marked confirmed (68.3%), 17 ambiguous (28.3%), and two not confirmed (3.3%). Twenty-nine records describe playable characters (48.3%), while 31 describe non-playable characters (51.7%). Six of the 22 game-system records describe romance options that do not depend on gender (27.3%). We calculated these values during the dataset review. They show that the fields can support these questions, but they do not show that Press Q or Quiu can already return checked percentages automatically (PRESS Q DATASET, 2026; PRESS Q CODE, 2026).")
    doc.add_paragraph("Table 2 - Illustrative audit calculations and development requirements", style="Conference Caption")
    table2 = doc.add_table(rows=1, cols=4)
    headers2 = ["Question", "Matching / total records", "Test result", "What it means"]
    for idx, text in enumerate(headers2):
        format_cell(table2.rows[0].cells[idx], text, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table2.rows[0])
    rows2 = [
        ("Confirmed among current character records?", "41 / 60", "68.3%", "Possible for this snapshot; still needs automation and testing."),
        ("Playable among current character records?", "29 / 60", "48.3%", "Possible for this snapshot; still needs automation and testing."),
        ("Romance not limited by gender among system records?", "6 / 22", "27.3%", "Possible within the system records; still needs testing."),
        ("Queer characters among all video game characters?", "No reliable total", "Cannot answer", "The dataset cannot support an industry-wide claim."),
        ("Intersectional combinations within character records?", "Data are inconsistent", "Not reliable yet", "The fields, missing values, and comparison groups need revision."),
    ]
    for index, row_data in enumerate(rows2):
        row = table2.add_row()
        for idx, text in enumerate(row_data):
            format_cell(row.cells[idx], text, align=WD_ALIGN_PARAGRAPH.LEFT)
            if index % 2 == 1:
                set_cell_shading(row.cells[idx], PALE_ALT)
    set_table_geometry(table2, [2700, 2050, 1422, 2900])
    set_table_borders(table2)
    doc.add_paragraph("Source: Press Q snapshot (PRESS Q DATASET, 2026); limitation on industry-wide totals (LGBTQ VIDEO GAME ARCHIVE, 2025a).", style="Conference Source")
    add_body("Table 2 gives a provisional answer to RQ1. The developing dataset can support repeatable calculations when the question clearly states what is being counted, which records are included, the total group, the date of the data, and how missing values are handled. However, the automatic calculation service and its tests still need to be built. The data cannot estimate the percentage of all games or all video game characters that are queer because the dataset is incomplete, was not selected as a representative sample, and has no known industry-wide total (PRESS Q CODE, 2026; PRESS Q DATASET, 2026; LGBTQ VIDEO GAME ARCHIVE, 2025a).")

    add_heading("5.3 Preparing for intersectional analysis", 2)
    add_body("Intersectional statistics are planned, but they are not reliable yet. The character file has one field intended to mark whether intersectionality is present and another field for written details. The first field currently contains 53 “no” values, three “yes” values, and four values that instead name a category such as class, ethnicity, race, or religion. Press Q therefore cannot yet calculate a trustworthy percentage for intersectional representation. A “no” may mean that an identity is absent, that no source documented it, or simply that the record is incomplete (PRESS Q DATASET, 2026). Future versions should separate missing data from research findings, allow several identity dimensions when sources support them, record uncertainty and sources, and avoid guessing identity from appearance or from silence. This work could address an area that the Archive itself describes as incomplete: its public website emphasizes gender and sexuality but gives less systematic attention to race, class, disability, character role, and cultural differences (CRENSHAW, 1989; D'IGNAZIO; KLEIN, 2020; LGBTQ VIDEO GAME ARCHIVE, 2025a).")

    add_heading("5.4 Using the data does not remove AI errors", 2)
    add_body("Quiu receives the Press Q data and instructions not to make unsupported guesses. Research suggests that outside evidence can improve factual accuracy, but it does not make every answer correct (LEWIS et al., 2020; SHUSTER et al., 2021). In the current version, the language model still interprets the records and may also count them while writing its response. It can therefore miscount, leave out an important limit, or produce a convincing statement that the data do not support (PRESS Q CODE, 2026; NIST, 2024).")
    doc.add_paragraph("Table 3 - Current safeguards and what still needs testing", style="Conference Caption")
    table3 = doc.add_table(rows=1, cols=3)
    headers3 = ["Current safeguard", "How it helps", "What still needs to be done"]
    for idx, text in enumerate(headers3):
        format_cell(table3.rows[0].cells[idx], text, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table3.rows[0])
    rows3 = [
        ("All three datasets are included with the question", "Keeps the answer focused on Press Q data", "Use search or calculation tools as the dataset grows"),
        ("Characters, systems, and readings are kept separate", "Reduces category and counting errors", "Add automatic tests for each type of calculation"),
        ("Instructions explain missing data and comparison groups", "Encourages limited answers and appropriate refusals", "Test refusals with misleading and multilingual questions"),
        ("A low temperature setting is used", "May support more consistent wording", "Do not treat this setting as proof of accuracy"),
        ("Records contain sources and evidence notes", "Lets the interface show where information came from", "Cite the exact records and keep source links in each answer"),
    ]
    for index, row_data in enumerate(rows3):
        row = table3.add_row()
        for idx, text in enumerate(row_data):
            format_cell(row.cells[idx], text)
            if index % 2 == 1:
                set_cell_shading(row.cells[idx], PALE_ALT)
    set_table_geometry(table3, [2850, 2700, 3522])
    set_table_borders(table3)
    doc.add_paragraph("Source: architecture audit of Press Q (PRESS Q CODE, 2026), interpreted against grounded-generation research and AI risk guidance (LEWIS et al., 2020; SHUSTER et al., 2021; NIST, 2024).", style="Conference Source")

    add_heading("5.5 Which Archive-related gaps could Press Q address?", 2)
    add_body("Press Q could make some Archive-related data easier to compare and calculate, but it cannot make the source material complete. The Archive focuses on documenting history and helping people discover examples. Press Q focuses on separating record types, defining fields, showing the total used in a calculation, linking evidence, and making missing information visible (LGBTQ VIDEO GAME ARCHIVE, 2025a, 2025b; PRESS Q CODE, 2026; PRESS Q DATASET, 2026).")
    doc.add_paragraph("Table 4 - Gaps Press Q may address and remaining limits", style="Conference Caption")
    table4 = doc.add_table(rows=1, cols=3)
    headers4 = ["Research gap", "What Press Q may add", "Remaining limit"]
    for idx, text in enumerate(headers4):
        format_cell(table4.rows[0].cells[idx], text, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table4.rows[0])
    rows4 = [
        ("Posts and categories can overlap", "Uses separate records for characters, systems, and readings", "The structure helps, but shared identifiers are still needed"),
        ("No total for industry-wide percentages", "Shows the total used for local dataset calculations", "Cannot provide a total that does not exist"),
        ("Evidence comes in different forms", "Records sources, confirmation, confidence, notes, and opposing evidence", "Fields remain incomplete and need a coding guide"),
        ("Limited public intersectional coding", "Plans fields and calculations for several identity dimensions", "Early stage; current values are not statistically reliable"),
        ("Games change across versions and languages", "Records platform, version, and review date", "Platform/version is filled for 32 of 60 character records"),
        ("New contributions must be reviewed", "Offers a public draft form and an administration page", "Permanent storage and review rules are not finished"),
        ("Complex records can be difficult to search", "Plans visual filters and questions in everyday language", "Usefulness and accuracy have not been formally tested"),
    ]
    for index, row_data in enumerate(rows4):
        row = table4.add_row()
        for idx, text in enumerate(row_data):
            format_cell(row.cells[idx], text)
            if index % 2 == 1:
                set_cell_shading(row.cells[idx], PALE_ALT)
    set_table_geometry(table4, [2700, 3300, 3072])
    set_table_borders(table4)
    doc.add_paragraph("Source: comparison of the Archive’s public methodology with the Press Q code and dataset snapshots (LGBTQ VIDEO GAME ARCHIVE, 2025a, 2025b; PRESS Q CODE, 2026; PRESS Q DATASET, 2026).", style="Conference Source")

    add_heading("6 Discussion", 1)
    add_heading("6.1 Clear statistics are a development goal", 2)
    add_body("Press Q should not simply produce a percentage without explanation. Its value would be showing how the number was created. A reliable result should state what was counted, which filters were used, how many records matched, the total number examined, the dataset date, the missing values, the strength of the evidence, and whether several identity dimensions were combined. Research on data and visualization shows that calculations reflect documented research choices rather than neutral facts (DRUCKER, 2011; D'IGNAZIO; KLEIN, 2020; WILKINSON et al., 2016).")
    add_body("For example, 68.3% means that 41 of the 60 character records in the 22 August 2026 Press Q snapshot are marked confirmed. It does not prove that the current website can already calculate this value reliably. It also does not mean that 68.3% of characters in the Archive, in published games, or in the game industry are queer (PRESS Q CODE, 2026; PRESS Q DATASET, 2026; LGBTQ VIDEO GAME ARCHIVE, 2025a).")

    add_heading("6.2 Intersectionality in the planned data model", 2)
    add_body("RQ4 concerns future work. If the dataset later contains consistent, well-sourced, and sufficiently complete fields for several identity dimensions, Press Q could answer questions such as: what proportion of the documented queer character records also contain evidence about race, ethnicity, disability, class, religion, or nationality? The result would still need a clear total and separate breakdowns. One general “intersectionality percentage” could hide important differences between groups and experiences (CRENSHAW, 1989; D'IGNAZIO; KLEIN, 2020). This feature could extend an area that the Archive’s public website identifies as incomplete, but the current Press Q data cannot support it responsibly (LGBTQ VIDEO GAME ARCHIVE, 2025a; PRESS Q DATASET, 2026).")

    add_heading("6.3 Quiu should explain checked calculations", 2)
    add_body("RQ2 requires a distinction between using evidence and checking an answer. The dataset limits the information given to Quiu, and the instructions set useful boundaries. However, the language model still generates text; it is not a calculation engine (PRESS Q CODE, 2026; NIST, 2024). A separate tool should therefore read the requested record type and filters, check the values, perform the calculation, and return the supporting records. Quiu should explain that checked result in clear language (NIST, 2024).")
    add_body("This design would also make Quiu easier to evaluate. Test questions could compare its answer with the checked calculation. The evaluation could measure whether the matching and total counts are correct, whether limits are preserved, whether citations point to the right records, whether the system refuses unsupported questions, and whether it works consistently across languages. Research supports this general approach, but Press Q still needs testing with queer-game researchers and the communities represented in the data (SHUSTER et al., 2021; NIST, 2024).")

    add_heading("6.4 Press Q should complement, not correct, the Archive", 2)
    add_body("RQ3 has only a provisional answer. Press Q is being designed to make selected Archive material and new research easier to compare, calculate, and search. It can already record details that the Archive’s public website does not emphasize, such as confidence in the evidence, platform or version, early intersectionality notes, and research status. However, the calculation tools and data coverage remain incomplete (LGBTQ VIDEO GAME ARCHIVE, 2025a; PRESS Q CODE, 2026; PRESS Q DATASET, 2026). Press Q cannot recover games, languages, sources, versions, or industry totals that are missing from the available evidence. It also depends on the Archive’s historical and curatorial work (LGBTQ VIDEO GAME ARCHIVE, 2025a).")
    add_body("The relationship should therefore be respectful and two-way. Press Q should credit each record, link to Archive pages and original sources, share corrections when appropriate, and avoid using the Archive’s work without acknowledgement. The Archive specifically asks researchers to give credit because the project represents hundreds of hours of labor, much of it unpaid (LGBTQ VIDEO GAME ARCHIVE, 2025a).")

    add_heading("7 Limitations and Future Work", 1)
    add_body("The dataset is incomplete and focuses on a limited group of titles and years. The records were selected for the project rather than drawn as a representative sample. All percentages in this paper were calculated during our review; they are not checked outputs from the current application (PRESS Q CODE, 2026; PRESS Q DATASET, 2026). We did not verify every source claim, review the Archive’s full master list, interview the Archive team, or compare decisions made by independent researchers coding the same records (LGBTQ VIDEO GAME ARCHIVE, 2025a; PRESS Q DATASET, 2026).")
    add_body("The intersectionality fields do not consistently distinguish between an identity that is absent, an identity that is unknown, information that no source documented, and a finding that researchers explicitly checked. The platform/version field is filled in for only 32 of the 60 character records. Missing or inconsistent values can change search results and calculations. These fields need revision before broader analysis (PRESS Q DATASET, 2026).")
    add_body("Quiu has not yet received a published evaluation of factual accuracy, calculations, citations, ease of use, accessibility, multilingual performance, or responses to misleading questions. Giving the model every record may also become impractical as the dataset grows. Future versions will need tools that find relevant records and perform calculations separately from the language model (PRESS Q CODE, 2026; LEWIS et al., 2020; NIST, 2024).")
    add_body("Future work should move the records to a permanent relational database with stable identifiers and a history of changes. The project should publish a coding guide that explains intersectionality and separates “unknown,” “not documented,” “not applicable,” and “none found.” It should record several identity dimensions only when sources support them, compare work by independent coders, include community review, build checked calculation tools, cite the exact records used in each answer, test supported and unsupported questions in several languages, and publish versioned data and code when rights allow (CRENSHAW, 1989; WILKINSON et al., 2016; NIST, 2024).")

    add_heading("8 Conclusion", 1)
    add_body("Press Q is an ongoing project that explores how structured evidence about queer games could support questions that are difficult to answer through a browsing-focused archive. The current review shows that some dataset-specific percentages, such as confirmed or playable character records, are possible because the records define what is being counted. However, the automatic calculation service is not finished. Questions about the whole game industry also remain impossible because the available data are incomplete and there is no reliable industry total (PRESS Q CODE, 2026; PRESS Q DATASET, 2026; LGBTQ VIDEO GAME ARCHIVE, 2025a).")
    add_body("Quiu makes the records easier to explore through everyday language. The dataset gives the model evidence to use; it does not train the model, and it cannot prevent every unsupported answer. The proposed design therefore separates the tasks: a calculation tool should produce checked numbers, Quiu should explain them, each answer should show its sources, and the limits of the data should remain visible (PRESS Q CODE, 2026; LEWIS et al., 2020; SHUSTER et al., 2021; NIST, 2024).")
    add_body("Press Q may help address some gaps around the LGBTQ Video Game Archive, but it remains dependent on the Archive’s historical work. One important future contribution is the ability to calculate clear intersectional results within the Press Q dataset—an area that the Archive’s public website does not emphasize. The project’s research value will depend on expanding and checking the data, improving the intersectionality fields without guessing identity, testing Quiu with real users, and preserving uncertainty, credit, and the social context of queer archival knowledge (LGBTQ VIDEO GAME ARCHIVE, 2025a; CRENSHAW, 1989; D'IGNAZIO; KLEIN, 2020; PRESS Q DATASET, 2026).")

    add_heading("Declarations for Conference Submission", 1)
    declarations = [
        ("Author information", "Removed for review; restore names, affiliations, ORCIDs, and correspondence details according to the conference policy."),
        ("Funding", "[State the funding source and grant number, or declare that no specific funding was received.]"),
        ("Conflicts of interest", "[State any conflicts, or declare that the authors have no conflicts of interest.]"),
        ("Data and code availability", "The cited Press Q code and dataset snapshots are anonymized for review. Add the public repository URL and a versioned DOI before camera-ready publication."),
        ("Generative AI use", "Suggested disclosure: OpenAI Codex assisted with source review, drafting, translation, and document formatting. The human authors are responsible for checking every source, calculation, interpretation, and final claim under the conference policy."),
    ]
    for label, value in declarations:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(3)
        lead = p.add_run(label + ": ")
        set_run_font(lead, size=9.8, bold=True, color=ACCENT_DARK)
        body = p.add_run(value)
        set_run_font(body, size=9.8, color=INK)

    add_heading("References", 1)
    references = [
        ("CRENSHAW, Kimberlé. Demarginalizing the intersection of race and sex: a Black feminist critique of antidiscrimination doctrine, feminist theory and antiracist politics. University of Chicago Legal Forum, v. 1989, n. 1, p. 139–167, 1989. ", "Institutional text", "https://chicagounbound.uchicago.edu/uclf/vol1989/iss1/8/"),
        ("D'IGNAZIO, Catherine; KLEIN, Lauren F. Data Feminism. Cambridge: MIT Press, 2020. ", "DOI", "https://doi.org/10.7551/mitpress/11805.001.0001"),
        ("DRUCKER, Johanna. Humanities approaches to graphical display. Digital Humanities Quarterly, v. 5, n. 1, 2011. ", "Full text", "https://www.digitalhumanities.org/dhq/vol/5/1/000091/000091.html"),
        ("LEWIS, Patrick et al. Retrieval-augmented generation for knowledge-intensive NLP tasks. Advances in Neural Information Processing Systems, v. 33, p. 9459–9474, 2020. ", "Preprint", "https://arxiv.org/abs/2005.11401"),
        ("LGBTQ VIDEO GAME ARCHIVE. About (please read first!). Methodology and limitations updated 4 June 2025. Accessed 22 August 2026. ", "Archive methodology", "https://lgbtqgamearchive.com/about/about-archive/"),
        ("LGBTQ VIDEO GAME ARCHIVE. Category descriptions. Accessed 22 August 2026. ", "Category descriptions", "https://lgbtqgamearchive.com/resources/category-descriptions/"),
        ("NATIONAL INSTITUTE OF STANDARDS AND TECHNOLOGY (NIST). Artificial Intelligence Risk Management Framework: Generative Artificial Intelligence Profile. NIST AI 600-1. Gaithersburg, 2024. ", "Official PDF", "https://nvlpubs.nist.gov/nistpubs/ai/NIST.AI.600-1.pdf"),
        ("PRESS Q CODE. Press Q and Quiu source-code snapshot inspected 22 August 2026. Anonymized local repository for conference review. 2026.", None, None),
        ("PRESS Q DATASET. pressq_seed_dataset.csv; game_queer_systems.csv; queer_readings.csv. Snapshot dated 22 August 2026. Anonymized for conference review. 2026.", None, None),
        ("RUBERG, Bonnie. Video Games Have Always Been Queer. New York: New York University Press, 2019. ", "Publisher page", "https://nyupress.org/9781479831036/video-games-have-always-been-queer/"),
        ("RUBERG, Bonnie; SHAW, Adrienne (eds.). Queer Game Studies. Minneapolis: University of Minnesota Press, 2017. ", "Publisher page", "https://www.upress.umn.edu/9781517900373/queer-game-studies/"),
        ("SHAW, Adrienne. Gaming at the Edge: Sexuality and Gender at the Margins of Gamer Culture. Minneapolis: University of Minnesota Press, 2014. ", "Publisher page", "https://www.upress.umn.edu/9780816693160/gaming-at-the-edge/"),
        ("SHAW, Adrienne; LAUTERIA, Evan W.; YANG, Hannah; PERSAUD, Christopher J.; COLE, Alayna M. Counting queerness in games: trends in LGBTQ digital game representation, 1985–2005. International Journal of Communication, v. 13, p. 1544–1569, 2019. ", "Full text", "https://ijoc.org/index.php/ijoc/article/view/9754/2611"),
        ("SHAW, Adrienne; PERSAUD, Christopher J. Beyond texts: using queer readings to document LGBTQ game content. First Monday, v. 25, n. 8, 2020. ", "DOI", "https://doi.org/10.5210/fm.v25i8.10439"),
        ("SHUSTER, Kurt; POFF, Spencer; CHEN, Moya; KIELA, Douwe; WESTON, Jason. Retrieval augmentation reduces hallucination in conversation. Findings of EMNLP 2021, p. 3784–3803, 2021. ", "DOI", "https://doi.org/10.18653/v1/2021.findings-emnlp.320"),
        ("WILKINSON, Mark D. et al. The FAIR Guiding Principles for scientific data management and stewardship. Scientific Data, v. 3, art. 160018, 2016. ", "DOI", "https://doi.org/10.1038/sdata.2016.18"),
    ]
    for text, label, url in references:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(-0.7)
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        set_run_font(r, size=9.0, color=INK)
        if label and url:
            add_hyperlink(p, label, url)

    for paragraph in doc.paragraphs:
        p_pr = paragraph._p.get_or_add_pPr()
        widow = p_pr.find(qn("w:widowControl"))
        if widow is None:
            widow = OxmlElement("w:widowControl")
            widow.set(qn("w:val"), "true")
            p_pr.append(widow)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()

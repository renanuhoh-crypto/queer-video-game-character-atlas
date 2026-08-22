from __future__ import annotations

from io import BytesIO
from pathlib import Path
import textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "artigo_cientifico_pressq.docx"

FONT = "Times New Roman"
INK = "202124"
MUTED = "5F6368"
ACCENT = "3B4C8A"
ACCENT_DARK = "29345F"
PALE = "EEF1F8"
PALE_ALT = "F7F8FC"
GRID = "B9C0D5"
WHITE = "FFFFFF"
CONTENT_DXA = 9072  # A4 with 2.5 cm left/right margins.


def set_run_font(run, size: float | None = None, bold=None, italic=None, color: str | None = None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def paragraph_border_bottom(paragraph, color="D9DDE8", size="6", space="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="5"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "true")
        tr_pr.append(cant_split)


def set_table_geometry(table, widths: list[int], indent=120):
    if sum(widths) != CONTENT_DXA:
        raise ValueError(f"Column widths must sum to {CONTENT_DXA}, got {sum(widths)}")
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        set_row_cant_split(row)
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Twips(widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def format_cell(cell, text: str, header=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.right_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    set_run_font(run, size=9.0 if not header else 9.2, bold=header, color=WHITE if header else INK)
    if header:
        set_cell_shading(cell, ACCENT_DARK)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_pr.append(r_fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    r_pr.append(size)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


def make_diagram() -> BytesIO:
    width, height = 1800, 900
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)

    def font(size, bold=False):
        candidates = [
            Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
            Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
        ]
        for candidate in candidates:
            if candidate.exists():
                return ImageFont.truetype(str(candidate), size=size)
        return ImageFont.load_default()

    title_font = font(34, True)
    body_font = font(27, False)
    small_font = font(22, False)
    dark = "#29345F"
    blue = "#E8EDF9"
    purple = "#F0E9F7"
    green = "#E9F4F0"

    def box(coords, fill, title, body):
        draw.rounded_rectangle(coords, radius=22, fill=fill, outline=dark, width=3)
        x1, y1, x2, y2 = coords
        draw.text(((x1 + x2) / 2, y1 + 34), title, font=title_font, fill=dark, anchor="ma")
        wrapped = "\n".join(textwrap.fill(line, width=24) for line in body.splitlines())
        draw.multiline_text(((x1 + x2) / 2, y1 + 95), wrapped, font=body_font, fill="#343A48", anchor="ma", align="center", spacing=8)

    def arrow(x1, y1, x2, y2):
        draw.line((x1, y1, x2 - 22, y2), fill=dark, width=5)
        draw.polygon([(x2 - 26, y2 - 13), (x2, y2), (x2 - 26, y2 + 13)], fill=dark)

    box((40, 45, 500, 245), blue, "PERSONAGENS", "identidades e codificações sustentadas por evidência")
    box((40, 310, 500, 510), purple, "SISTEMAS", "possibilidades de criação, romance, relação e família")
    box((40, 575, 500, 775), green, "LEITURAS QUEER", "interpretações, disputas, recepção e contraevidência")

    box((665, 220, 1170, 600), "#F7F8FC", "CAMADA ESTRUTURADA", "proveniência\nstatus de pesquisa\nconfiança da evidência\nplataforma e versão\ndata de revisão")
    box((1330, 270, 1760, 550), "#EEF1F8", "INTERFACE", "análises visuais\nconsulta em linguagem natural\nqualificações explícitas")

    arrow(500, 145, 665, 330)
    arrow(500, 410, 665, 410)
    arrow(500, 675, 665, 490)
    arrow(1170, 410, 1330, 410)
    draw.text((900, 850), "A separação impede que possibilidade lúdica ou interpretação seja contada como identidade canônica.", font=small_font, fill="#5F6368", anchor="mm")

    stream = BytesIO()
    image.save(stream, format="PNG", dpi=(220, 220))
    stream.seek(0)
    return stream


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    section.different_first_page_header_footer = False

    settings = doc.settings._element
    even_odd = settings.find(qn("w:evenAndOddHeaders"))
    if even_odd is None:
        even_odd = OxmlElement("w:evenAndOddHeaders")
        settings.append(even_odd)
    even_odd.set(qn("w:val"), "true")

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(0.75)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, before, after in (
        ("Heading 1", 13.0, 14, 6),
        ("Heading 2", 11.8, 11, 4),
        ("Heading 3", 11.0, 8, 3),
    ):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(ACCENT_DARK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.line_spacing = 1.05

    if "Abstract" not in styles:
        abstract_style = styles.add_style("Abstract", WD_STYLE_TYPE.PARAGRAPH)
    else:
        abstract_style = styles["Abstract"]
    abstract_style.font.name = FONT
    abstract_style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    abstract_style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    abstract_style.font.size = Pt(10)
    abstract_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abstract_style.paragraph_format.first_line_indent = Cm(0)
    abstract_style.paragraph_format.space_after = Pt(5)
    abstract_style.paragraph_format.line_spacing = 1.05

    if "Table Caption" not in styles:
        caption_style = styles.add_style("Table Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption_style = styles["Table Caption"]
    caption_style.font.name = FONT
    caption_style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    caption_style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    caption_style.font.size = Pt(9.5)
    caption_style.font.bold = True
    caption_style.font.color.rgb = RGBColor.from_string(ACCENT_DARK)
    caption_style.paragraph_format.first_line_indent = Cm(0)
    caption_style.paragraph_format.space_before = Pt(7)
    caption_style.paragraph_format.space_after = Pt(4)
    caption_style.paragraph_format.keep_with_next = True

    if "Table Source" not in styles:
        source_style = styles.add_style("Table Source", WD_STYLE_TYPE.PARAGRAPH)
    else:
        source_style = styles["Table Source"]
    source_style.font.name = FONT
    source_style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    source_style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    source_style.font.size = Pt(8.5)
    source_style.font.color.rgb = RGBColor.from_string(MUTED)
    source_style.paragraph_format.first_line_indent = Cm(0)
    source_style.paragraph_format.space_before = Pt(4)
    source_style.paragraph_format.space_after = Pt(6)
    source_style.paragraph_format.line_spacing = 1.0

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("PRESS Q | MANUSCRITO CIENTÍFICO")
    set_run_font(hr, size=8.5, bold=True, color=MUTED)

    even_header = section.even_page_header
    ehp = even_header.paragraphs[0]
    ehp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ehp.paragraph_format.space_after = Pt(0)
    ehr = ehp.add_run("PRESS Q | MANUSCRITO CIENTÍFICO")
    set_run_font(ehr, size=8.5, bold=True, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)

    even_footer = section.even_page_footer
    efp = even_footer.paragraphs[0]
    add_page_number(efp)

    doc.core_properties.title = "Do arquivo curado ao corpus estruturado"
    doc.core_properties.subject = "Limites do LGBTQ Video Game Archive e desenho metodológico do Press Q"
    doc.core_properties.keywords = "videogames; LGBTQ+; arquivos digitais; humanidades digitais; inteligência artificial"

    # First-page title block.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run("DO ARQUIVO CURADO AO CORPUS ESTRUTURADO: LIMITES DO LGBTQ VIDEO GAME ARCHIVE E O DESENHO METODOLÓGICO DO PRESS Q")
    set_run_font(r, size=16, bold=True, color=ACCENT_DARK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("From curated archive to structured corpus: limitations of the LGBTQ Video Game Archive and the methodological design of Press Q")
    set_run_font(r, size=11.5, italic=True, color=MUTED)

    for text, bold in (
        ("[AUTOR/A - PREENCHER ANTES DA SUBMISSÃO]", True),
        ("[Afiliação institucional | ORCID | e-mail para correspondência]", False),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(2 if bold else 12)
        r = p.add_run(text)
        set_run_font(r, size=10.5, bold=bold, color=INK if bold else MUTED)

    rule = doc.add_paragraph()
    rule.paragraph_format.first_line_indent = Cm(0)
    rule.paragraph_format.space_after = Pt(8)
    paragraph_border_bottom(rule, color="AEB7D2", size="8", space="4")

    def add_abstract(label: str, text: str):
        p = doc.add_paragraph(style="Abstract")
        lead = p.add_run(label + ": ")
        set_run_font(lead, size=10, bold=True, color=ACCENT_DARK)
        body = p.add_run(text)
        set_run_font(body, size=10, color=INK)

    add_abstract(
        "Resumo",
        "Este artigo examina os limites documentais, classificatórios e computacionais do LGBTQ Video Game Archive e analisa o Press Q como uma camada complementar de estruturação e consulta. Adotou-se um estudo de caso metodológico, combinando análise documental das páginas públicas e publicações associadas ao arquivo com auditoria descritiva do repositório e dos três conjuntos de dados do Press Q em 22 de agosto de 2026. O corpus do protótipo contém 94 registros: 60 personagens, 22 sistemas de jogo e 12 leituras queer. Os resultados identificam cinco grupos principais de limitações no arquivo de referência: cobertura não exaustiva e ausência de denominador; sobreposição entre categorias e unidades de análise; heterogeneidade das evidências; baixa operacionalização pública da interseccionalidade e da variação entre versões; e dependência de trabalho curatorial intensivo. O Press Q responde parcialmente por meio da separação entre personagens, possibilidades lúdicas e leituras de recepção, além de metadados de proveniência, revisão e confiança e de uma assistente conversacional restrita ao corpus. A auditoria, contudo, encontrou cobertura seletiva, metadados incompletos nos registros de personagens, inconsistência no campo de interseccionalidade e ausência de avaliação entre codificadores e usuários. Conclui-se que a estruturação computacional pode ampliar a rastreabilidade e a capacidade de consulta, mas não elimina a interpretação, a incompletude nem o trabalho político do arquivo.",
    )
    add_abstract("Palavras-chave", "videogames; representação LGBTQ+; arquivos digitais; humanidades digitais; dados situados; inteligência artificial responsável")
    add_abstract(
        "Abstract",
        "This article examines the documentary, classificatory, and computational limitations of the LGBTQ Video Game Archive and analyzes Press Q as a complementary layer for structuring and querying research information. The study combines documentary analysis of the archive's public pages and related publications with a descriptive audit of the Press Q repository and its three datasets on 22 August 2026. The prototype corpus contains 94 records: 60 characters, 22 game systems, and 12 queer readings. Five main limitation groups were identified: non-exhaustive coverage and the absence of a denominator; overlap between categories and units of analysis; heterogeneous evidence; limited public operationalization of intersectionality and version differences; and reliance on intensive curatorial labor. Press Q responds in part by separating characters, ludic possibilities, and reception histories, adding provenance and review metadata, and constraining a conversational assistant to the corpus. The audit nevertheless found selective coverage, incomplete character metadata, an inconsistent intersectionality flag, and no inter-coder or user evaluation. Computational structuring can therefore improve traceability and queryability, but it cannot eliminate interpretation, incompleteness, or the political work of archival practice.",
    )
    add_abstract("Keywords", "video games; LGBTQ+ representation; digital archives; digital humanities; situated data; responsible artificial intelligence")

    def add_heading(text: str, level=1):
        return doc.add_paragraph(text, style=f"Heading {level}")

    def add_body(text: str, indent=True, keep=False):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.75) if indent else Cm(0)
        p.paragraph_format.keep_with_next = keep
        r = p.add_run(text)
        set_run_font(r, size=11, color=INK)
        return p

    add_heading("1 Introdução", 1)
    add_body("Arquivos digitais dedicados à história LGBTQ+ dos videogames enfrentam uma tensão constitutiva. Para tornar experiências dispersas pesquisáveis, precisam nomear, agrupar e contar conteúdos que nem sempre são estáveis, explícitos ou comparáveis. Ao mesmo tempo, cada categoria produz fronteiras: decide o que será visível, o que permanecerá ambíguo e quais formas de evidência poderão sustentar uma classificação. Em estudos queer de jogos, essa tensão é especialmente importante porque identidade, prática de jogo, autoria, recepção e interpretação não coincidem necessariamente (SHAW, 2014; RUBERG; SHAW, 2017).")
    add_body("O LGBTQ Video Game Archive, coordenado por Adrienne Shaw, é uma infraestrutura pioneira para documentar conteúdos LGBTQ+ em jogos digitais desde a década de 1980. O projeto reúne páginas de jogos e sínteses baseadas em textos dos próprios jogos, declarações de produtores, jornalismo, wikis, vídeos, fóruns, críticas e leituras de fãs. Sua relevância decorre tanto da escala do levantamento quanto da decisão de preservar formas explícitas, implícitas, ambientais e queermente lidas de representação. Pesquisas derivadas do arquivo demonstraram o valor de combinar análise histórica quantitativa e atenção às práticas de recepção (SHAW et al., 2019; SHAW; PERSAUD, 2020).")
    add_body("O próprio arquivo, contudo, apresenta-se como coleção curada de informação e não como repositório tradicional de fontes primárias. Também reconhece que sua cobertura é incompleta, que as categorias podem ser debatidas e que não existe um denominador confiável para estimar a proporção de todos os jogos que possuem conteúdo LGBTQ+ (LGBTQ VIDEO GAME ARCHIVE, 2025). Essas ressalvas não diminuem sua contribuição. Ao contrário, fornecem o ponto de partida para examinar como limitações epistemológicas e infraestruturais poderiam ser tornadas mais explícitas em um modelo de dados complementar.")
    add_body("Este artigo pergunta: quais limitações do LGBTQ Video Game Archive se tornam relevantes quando suas informações são convertidas em um corpus estruturado, visualizável e consultável por inteligência artificial? O objetivo é analisar criticamente essas limitações e avaliar, em uma implementação inicial, como o Press Q procura respondê-las sem reivindicar substituição ou superioridade. A contribuição proposta é dupla: formular uma tipologia de limites do arquivo de referência e apresentar a separação entre personagens, sistemas de jogo e leituras queer como mecanismo de tipagem epistemológica. A hipótese é que separar o tipo de afirmação antes de contar reduz a conversão indevida de possibilidades lúdicas ou interpretações em identidades canônicas.")

    add_heading("2 Referencial teórico", 1)
    add_heading("2.1 Representação, leitura queer e arquivo", 2)
    add_body("A pesquisa sobre videogames e diferença não pode restringir-se à presença ou ausência de identidades. Shaw (2014) argumenta que jogadores vivenciam raça, gênero e sexualidade de maneira simultânea, enquanto Ruberg e Shaw (2017) ampliam o campo para práticas, afetos, corpos, regras e culturas de jogo. Ruberg (2019), por sua vez, propõe compreender a própria forma dos videogames por meio da teoria queer, deslocando a análise de uma busca exclusiva por personagens positivos ou explicitamente identificados. Esse campo torna insuficiente uma ontologia simples na qual cada conteúdo ocupa uma categoria fixa e autoevidente.")
    add_body("O trabalho do LGBTQ Video Game Archive enfrenta essa complexidade de modo produtivo. Shaw e Persaud (2020) defendem que leituras de fãs, textos do jogo, declarações de produtores e recepção crítica devem ser colocados em relação para preservar a história de como a queerness é produzida e reconhecida. A leitura queer não equivale à descoberta de uma identidade secreta e verdadeira; ela constitui uma prática histórica de interpretação. Por isso, apagá-la empobrece o arquivo, mas registrá-la no mesmo nível de uma identidade explicitamente afirmada também pode gerar confusão analítica.")
    add_body("A distinção entre canon, possibilidade e recepção é ainda mais delicada em jogos, cuja interatividade permite resultados mutuamente exclusivos. Um sistema pode autorizar romances com diferentes gêneros sem definir uma orientação fixa para cada personagem; um avatar configurável pode materializar possibilidades queer sem constituir uma personagem canônica; e uma leitura comunitária pode ser historicamente relevante mesmo quando contestada por criadores. A unidade de análise, portanto, antecede a contagem.")

    add_heading("2.2 Dados situados, interseccionalidade e visualização", 2)
    add_body("A crítica feminista dos dados ajuda a compreender por que a estruturação não é uma operação neutra. D'Ignazio e Klein (2020) situam dados e classificações dentro de relações de poder, perguntando quem conta, quem é contado e quais experiências desaparecem nos padrões de mensuração. Drucker (2011) propõe tratar dados humanísticos como construções interpretativas, e não como fatos simplesmente dados. Em um arquivo queer, essa perspectiva exige que categorias, lacunas, controvérsias e decisões curatoriais permaneçam visíveis na interface.")
    add_body("A interseccionalidade reforça esse requisito. Crenshaw (1989) mostrou que analisar eixos de poder como categorias isoladas pode apagar experiências constituídas justamente em suas interseções. Um registro que classifica sexualidade sem considerar raça, classe, deficiência, religião, nacionalidade ou contexto histórico não é necessariamente incorreto, mas oferece uma visão parcial. Incluir esses eixos, por outro lado, requer fontes e governança cuidadosas para impedir inferências baseadas em aparência, nome ou estereótipo.")
    add_body("Visualizações não resolvem essa tensão automaticamente. Gráficos podem conferir aparência de objetividade a um corpus seletivo e fazer percentuais internos parecerem estimativas do universo dos jogos. Uma interface humanística precisa declarar o denominador, mostrar sobreposições e distinguir 'desconhecido', 'não registrado' e 'nenhum marcador documentado'. No Press Q, essas distinções são tratadas como propriedades do processo de pesquisa, e não como identidades presumidas.")

    add_heading("2.3 Proveniência de dados e IA responsável", 2)
    add_body("Os princípios FAIR enfatizam que objetos digitais científicos devem ser localizáveis, acessíveis, interoperáveis e reutilizáveis, com metadados suficientes para sustentar proveniência e reuso (WILKINSON et al., 2016). Em um corpus cultural, aderir a esse horizonte não significa publicar toda fonte sem restrições, mas documentar origem, idioma, versão, data de revisão e estado da pesquisa. Esses metadados permitem que uma afirmação seja reavaliada quando um jogo recebe nova localização, atualização, expansão ou remasterização.")
    add_body("A introdução de inteligência artificial acrescenta outro nível de risco. O NIST AI Risk Management Framework recomenda governança, mapeamento de contexto, avaliação e gestão contínua dos riscos, além de supervisão humana (TABASSI, 2023). Uma assistente de pesquisa pode reduzir barreiras de acesso ao corpus, mas também produzir sínteses excessivamente confiantes. Limitar o contexto de resposta, explicitar ausências e preservar contraevidências são salvaguardas relevantes, porém sua eficácia precisa ser testada empiricamente.")

    add_heading("3 Metodologia", 1)
    add_heading("3.1 Desenho da pesquisa", 2)
    add_body("Foi realizado um estudo de caso exploratório de natureza metodológica, combinando análise documental qualitativa e auditoria descritiva de dados. O caso primário é a relação entre o LGBTQ Video Game Archive e o Press Q, protótipo de arquivo de humanidades digitais que reutiliza e reorganiza parte do conhecimento produzido por fontes arquivísticas, acadêmicas, jornalísticas e comunitárias. O estudo não avalia a exatidão de todas as entradas do arquivo de referência e não pretende produzir uma comparação de desempenho entre plataformas.")
    add_body("A análise documental considerou a página metodológica 'About (please read first!)', as descrições de categorias, a página de personagens, a listagem de tipos de conteúdo e as publicações de Shaw et al. (2019) e Shaw e Persaud (2020). Foram codificados cinco eixos: cobertura e denominador; unidade de análise e contagem; evidência e proveniência; interseccionalidade, versão e localização; e sustentabilidade do trabalho curatorial. A crítica distingue limitações declaradas pelo próprio arquivo de inferências derivadas de sua arquitetura pública.")

    add_heading("3.2 Materiais e unidade de observação", 2)
    add_body("A auditoria do Press Q utilizou uma captura do repositório de desenvolvimento em 22 de agosto de 2026. Foram examinados o código da interface, a rota da assistente Quiu, a documentação metodológica e três arquivos CSV: pressq_seed_dataset.csv, game_queer_systems.csv e queer_readings.csv. A unidade de observação variou conforme a tabela: personagem identificável, affordance ou sistema no nível do jogo e leitura crítica ou comunitária documentada.")
    add_body("Os registros de personagens contêm 28 campos, incluindo jogo, papel narrativo, jogabilidade, gênero, sexualidade, confirmação, situação queer, interseccionalidade, fonte, versão e revisão. Sistemas e leituras possuem 17 campos cada. Um registro de sistema descreve o que o jogo permite ao jogador configurar ou realizar; um registro de leitura documenta a existência e o contexto de uma interpretação. Nenhuma dessas unidades deve ser convertida automaticamente em identidade de personagem.")

    add_heading("3.3 Procedimentos analíticos", 2)
    add_body("Foram calculadas frequências absolutas e percentuais internos a cada unidade. Para os registros de personagens, examinaram-se ano, jogabilidade, gênero registrado, sexualidade registrada, forma de confirmação, situação queer e completude de metadados. Para sistemas, analisaram-se tipo, dependência do jogador, disponibilidade e confiança da evidência. Para leituras, analisaram-se tipo, situação da leitura e confiança. Campos multivalorados não foram somados como categorias mutuamente exclusivas. Não foram aplicados testes inferenciais porque o corpus é curado, pequeno e não probabilístico.")
    add_body("A completude foi definida como presença de valor não vazio, sem inferir a qualidade substantiva do conteúdo. A auditoria também verificou aderência aos valores controlados. Percentuais referem-se exclusivamente ao corpus do Press Q na data da captura. Títulos de jogos e séries foram tratados como rótulos textuais, de modo que a contagem de títulos distintos não equivale necessariamente a franquias, edições ou obras intelectuais únicas.")

    add_heading("3.4 Considerações éticas", 2)
    add_body("A pesquisa analisou registros documentais e código do projeto, sem recrutamento de participantes. Ainda assim, a classificação de identidades exige cautela: o estudo não infere gênero, sexualidade, raça, etnia, religião ou deficiência a partir de aparência, comportamento ou nome. Leituras queer são preservadas como recepção, inclusive quando contestadas, mas não alimentam totais de identidades confirmadas. Antes da submissão, os autores devem verificar as exigências específicas do periódico e da instituição quanto a ética em pesquisa, direitos autorais, imagens, contribuições comunitárias e divulgação do uso de IA.")

    add_heading("4 Resultados", 1)
    add_heading("4.1 Limitações documentais e infraestruturais do LGBTQ Video Game Archive", 2)
    add_heading("4.1.1 Estatuto arquivístico, cobertura e denominador", 3)
    add_body("A primeira limitação é explicitamente reconhecida: o site não se apresenta como arquivo tradicional de fontes primárias, mas como coleção curada de informação sobre conteúdos LGBTQ+ e queermente lidos. Parte dos materiais primários foi destinada a outra instituição, enquanto o site funciona como porta de entrada para pesquisa. Essa distinção afeta preservação, citação e verificabilidade: a síntese é valiosa, mas depende da continuidade dos links e das fontes externas às quais remete.")
    add_body("A cobertura também é incompleta por desenho. Em atualização de 4 de junho de 2025, o projeto informou possuir mais de 1.200 jogos em sua lista-mestra, mas pesquisa concluída para cerca de 400. O arquivo não conhece o número total confiável de jogos publicados por ano nem pode garantir que todos os casos LGBTQ+ tenham sido identificados. Assim, aumentos no número de registros não podem ser convertidos diretamente em prevalência relativa na indústria (LGBTQ VIDEO GAME ARCHIVE, 2025). Trata-se de viés de seleção e sobrevivência documental: obras recentes, acessíveis, anglófonas, discutidas online ou presentes em listas anteriores têm maior probabilidade de serem encontradas.")

    add_heading("4.1.2 Sobreposição categorial e instabilidade da contagem", 3)
    add_body("A arquitetura pública utiliza páginas de jogos, posts de conteúdos e categorias WordPress. Um mesmo post pode receber várias categorias, aparecer em diferentes listas e reunir mais de um exemplo. A taxonomia combina entidades de naturezas distintas: personagens, ações, artefatos, localizações, relacionamentos, modificações, homofobia ou transfobia, mudanças de localização e jogos inteiramente queer (LGBTQ VIDEO GAME ARCHIVE, s.d.-a). Essa riqueza favorece navegação temática, mas dificulta interpretar contagens como registros independentes.")
    add_body("O problema não é a sobreposição em si, que é adequada a fenômenos queer e interseccionais, mas a ausência de uma unidade analítica uniforme em cada número apresentado pela interface. Uma categoria pode contar posts, personagens, momentos, jogos ou combinações desses elementos. Sem identificadores persistentes e uma tabela de relações, somas entre categorias produzem duplicação. Para análise quantitativa, torna-se necessário reconstruir manualmente o que cada item representa.")

    add_heading("4.1.3 Evidência heterogênea, canon e recepção", 3)
    add_body("O arquivo mobiliza fontes heterogêneas: texto do jogo, declarações oficiais, jornalismo, wikis, vídeos, fóruns, críticas e interpretações de fãs. Essa abertura é metodologicamente importante para preservar conteúdos pouco documentados, mas as fontes oferecem graus diferentes de proximidade, estabilidade e autoridade. A distinção público-visível entre explícito, implícito e queermente lido reduz o problema, porém não o elimina. Evidência da existência de uma leitura não é a mesma coisa que evidência da identidade atribuída pela leitura.")
    add_body("Shaw e Persaud (2020) defendem corretamente que práticas de recepção fazem parte da história dos textos e não devem ser descartadas. A limitação surge quando a infraestrutura precisa simultaneamente preservar a interpretação e impedir que ela seja retomada, por mecanismos de busca ou análises posteriores, como fato canônico. Isso exige campos separados para resumo da leitura, situação, contraevidência, resposta de criadores e confiança na documentação.")

    add_heading("4.1.4 Interseccionalidade, versões e localização", 3)
    add_body("O arquivo informa que a codificação pública do site se concentra em gênero e sexualidade. A planilha subjacente busca acompanhar outras dimensões, mas é disponibilizada mediante contato, não como parte integrada da navegação pública (LGBTQ VIDEO GAME ARCHIVE, 2025). Como consequência, raça, classe, deficiência, religião e papel narrativo permanecem menos operacionalizados para consulta comparativa, embora a literatura demonstre que essas dimensões são simultâneas (SHAW, 2014).")
    add_body("O projeto reconhece mudanças de localização e mantém uma categoria específica para conteúdos alterados entre mercados. Ainda assim, categorias de site não substituem metadados por registro sobre plataforma, edição, patch, idioma e data da observação. Jogos são objetos mutáveis: diálogos, romances, opções de criação e censura podem variar. Sem versionamento granular, uma afirmação pode ser correta para uma edição e inadequada para outra.")

    add_heading("4.1.5 Trabalho, manutenção e preservação", 3)
    add_body("O arquivo declara que representa centenas de horas de trabalho, muitas vezes não remunerado, e sua página inicial registra capacidade limitada para receber colaboração voluntária. A pesquisa de um único jogo pode exigir localizar materiais efêmeros e conciliar versões conflitantes. Portanto, incompletude e demora não são apenas problemas técnicos: refletem recursos institucionais, cuidado curatorial e dependência de trabalho acadêmico e comunitário.")
    add_body("A mesma dependência aparece na preservação. Fóruns, vídeos, wikis e páginas pessoais podem desaparecer; links podem mudar; materiais podem ser removidos por direitos autorais ou encerramento de plataformas. O arquivo sintetiza e aponta para essas fontes, mas não controla sua continuidade. Uma infraestrutura complementar precisa registrar proveniência e data de acesso e, quando juridicamente permitido, adotar estratégias institucionais de preservação.")

    add_heading("4.2 O modelo complementar do Press Q", 2)
    add_body("O Press Q reorganiza o problema em três unidades explícitas. Personagens registram identidades ou codificações sustentadas por evidência no nível do indivíduo. Sistemas registram affordances que permitem configurar gênero, selecionar pronomes, estabelecer romances, casar ou formar famílias. Leituras queer registram interpretações críticas e comunitárias, com situação e contraevidência. A Figura 1 sintetiza o modelo.")

    caption = doc.add_paragraph("Figura 1 - Separação das unidades de análise e camada de consulta do Press Q", style="Table Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig_p = doc.add_paragraph()
    fig_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig_p.paragraph_format.first_line_indent = Cm(0)
    fig_p.paragraph_format.space_after = Pt(2)
    run = fig_p.add_run()
    picture = run.add_picture(make_diagram(), width=Cm(15.7))
    picture._inline.docPr.set("descr", "Diagrama com três unidades separadas - personagens, sistemas e leituras queer - conectadas a uma camada estruturada de proveniência e, depois, à interface de análises e consulta.")
    source = doc.add_paragraph("Fonte: elaboração própria a partir da metodologia do Press Q.", style="Table Source")
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_body("Essa separação funciona como tipagem epistemológica: cada linha declara que tipo de afirmação está sendo feita. A estrutura inclui fonte, idioma, origem da descoberta, situação da pesquisa, confiança, plataforma ou versão e última revisão. Nas leituras, há campos próprios para situação e contraevidência; nos sistemas, registram-se dependência do jogador, disponibilidade e limitações. Desse modo, a possibilidade de um romance não transforma todos os personagens compatíveis em bissexuais ou pansexuais, e uma leitura contestada não integra o total de identidades confirmadas.")
    add_body("A camada de interface oferece análises por jogabilidade, gênero, sexualidade, interseccionalidade, escala de produção, tipo de sistema, dependência do jogador, disponibilidade, situação da pesquisa e confiança. Quiu, a assistente conversacional, recebe os três conjuntos de dados como contexto e é instruída a responder no idioma da pergunta, declarar lacunas e manter as unidades separadas. No snapshot analisado, a implementação utilizava um modelo da OpenAI com temperatura 0,35. Trata-se de restrição por instrução e contexto, não de garantia formal contra erro.")

    add_heading("4.3 Composição e vieses do corpus inicial", 2)
    add_body("O corpus continha 94 registros e 38 rótulos distintos de jogos ou séries quando as três unidades eram combinadas. A Tabela 1 mostra que cada unidade possui denominador e recorte temporal próprios. Somar títulos por linha produziria duplicação, pois alguns jogos aparecem em mais de uma unidade.")

    cap = doc.add_paragraph("Tabela 1 - Composição do corpus do Press Q em 22 de agosto de 2026", style="Table Caption")
    table = doc.add_table(rows=1, cols=5)
    headers = ["Unidade", "Registros", "Títulos distintos", "Período", "Objeto registrado"]
    for idx, text in enumerate(headers):
        format_cell(table.rows[0].cells[idx], text, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0])
    rows = [
        ("Personagens", "60", "24", "1990-2020", "Personagem identificável e evidência de identidade ou codificação"),
        ("Sistemas", "22", "14", "1990-2018", "Possibilidade no nível do jogo, com escopo e dependência do jogador"),
        ("Leituras queer", "12", "7", "1984-2016", "Interpretação documentada, situação, contraevidência e recepção"),
    ]
    for row_data in rows:
        row = table.add_row()
        for idx, text in enumerate(row_data):
            align = WD_ALIGN_PARAGRAPH.CENTER if idx in (1, 2, 3) else WD_ALIGN_PARAGRAPH.LEFT
            format_cell(row.cells[idx], text, align=align)
    set_table_geometry(table, [1500, 1200, 1350, 1150, 3872])
    set_table_borders(table)
    doc.add_paragraph("Fonte: elaboração própria com base nos três arquivos CSV do Press Q. Títulos distintos são rótulos textuais e podem incluir séries.", style="Table Source")

    add_body("Entre os 60 personagens, 29 (48,3%) eram jogáveis e 31 (51,7%) não jogáveis. O campo de situação queer classificava 41 registros (68,3%) como confirmados, 17 (28,3%) como ambíguos e dois (3,3%) como não confirmados. A forma de confirmação variava: 19 registros (31,7%) eram explícitos no jogo; 15 (25,0%) dependiam de fonte externa; 13 (21,7%) eram não explícitos; nove (15,0%) eram ambíguos; e quatro (6,7%) tinham confirmação de desenvolvedores. Essas categorias descrevem o corpus, não a indústria.")
    add_body("A distribuição temporal era concentrada: 29 dos 60 personagens (48,3%) pertenciam a jogos lançados em 2014 ou 2015. Esse resultado evidencia a dependência de um conjunto-semente e impede leituras históricas de tendência. O campo de gênero registrava 36 homens, 13 mulheres, cinco pessoas não binárias, três homens trans, uma mulher trans, uma pessoa gênero-fluido e um caso desconhecido. Como o campo usa rótulos gerais e específicos em uma única coluna, análises devem evitar tratar homens trans como oposição a homens ou mulheres trans como oposição a mulheres. Uma revisão futura pode separar gênero geral e dimensão trans documentada.")
    add_body("Nos sistemas, 19 de 22 registros (86,4%) dependiam integralmente de ação do jogador e três (13,6%) dependiam parcialmente. Quatorze sistemas (63,6%) eram opcionais, cinco (22,7%) padrão, dois (9,1%) condicionais e um (4,5%) associado a expansão. Romance independente de gênero era o tipo mais frequente, com seis registros (27,3%). Esses resultados mostram por que sistemas não devem ser usados para inferir identidades fixas: a maior parte do conteúdo depende de escolhas e pode não ocorrer em uma partida específica.")
    add_body("Entre as 12 leituras queer, cinco eram contestadas, cinco eram classificadas como queermente lidas e duas tinham refutação de criadores. A confiança se referia à documentação da existência e do contexto da leitura, não à verdade canônica da identidade interpretada: seis registros tinham confiança baixa, quatro média e dois alta. Essa distinção preserva recepção histórica sem transformar confiança documental em probabilidade de a leitura estar correta.")

    add_heading("4.4 Completude de metadados e limites residuais", 2)
    add_body("A auditoria encontrou maior completude nos conjuntos de sistemas e leituras do que nos personagens. Em todo o corpus, 87 de 94 registros (92,6%) possuíam situação de pesquisa, confiança da evidência e data da última revisão; 93 (98,9%) tinham fonte; 86 (91,5%) registravam idioma e origem da descoberta; e 62 (66,0%) informavam plataforma ou versão. A Tabela 2 detalha a diferença entre unidades.")

    cap = doc.add_paragraph("Tabela 2 - Completude de metadados por unidade, n (%)", style="Table Caption")
    table2 = doc.add_table(rows=1, cols=4)
    headers2 = ["Campo", "Personagens (n=60)", "Sistemas (n=22)", "Leituras (n=12)"]
    for idx, text in enumerate(headers2):
        format_cell(table2.rows[0].cells[idx], text, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table2.rows[0])
    rows2 = [
        ("Fonte da evidência", "59 (98,3)", "22 (100,0)", "12 (100,0)"),
        ("Idioma da fonte", "52 (86,7)", "22 (100,0)", "12 (100,0)"),
        ("Origem da descoberta", "52 (86,7)", "22 (100,0)", "12 (100,0)"),
        ("Situação da pesquisa", "53 (88,3)", "22 (100,0)", "12 (100,0)"),
        ("Confiança da evidência", "53 (88,3)", "22 (100,0)", "12 (100,0)"),
        ("Plataforma ou versão", "32 (53,3)", "18 (81,8)", "12 (100,0)"),
        ("Última revisão", "53 (88,3)", "22 (100,0)", "12 (100,0)"),
    ]
    for row_data in rows2:
        row = table2.add_row()
        for idx, text in enumerate(row_data):
            format_cell(row.cells[idx], text, align=WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    set_table_geometry(table2, [2600, 2157, 2157, 2158])
    set_table_borders(table2)
    doc.add_paragraph("Fonte: elaboração própria. Completude indica apenas presença de valor não vazio, não validade ou qualidade da informação.", style="Table Source")

    add_body("A interseccionalidade revelou um problema de validação. O campo intersectionality_present deveria operar como indicador, mas continha 53 valores 'no', três valores 'yes' e quatro valores que nomeavam diretamente eixos como classe, etnia, raça ou religião. Assim, sete registros possuíam valor diferente de 'no', mas apenas três obedeciam ao formato binário esperado. A inconsistência pode alterar filtros e percentuais e mostra que um esquema mais rico não garante dados consistentes sem validação na entrada e revisão curatorial.")
    add_body("Também não havia, no snapshot, medida de concordância entre codificadores, registro público de adjudicação de conflitos ou estudo com usuários. A interface administra os dados por arquivos CSV, apropriados ao protótipo, mas inadequados para colaboração concorrente, histórico de versões e implantação serverless persistente. Por fim, Quiu não possuía uma avaliação sistemática publicada sobre fidelidade, preservação de qualificações, sensibilidade a idioma ou comportamento diante de perguntas adversariais.")

    add_heading("4.5 Matriz comparativa", 2)
    add_body("A Tabela 3 resume como o Press Q transforma limitações observadas em decisões de desenho e quais riscos permanecem. As respostas devem ser lidas como complementares: o Press Q depende do trabalho histórico e das fontes reunidas pelo LGBTQ Video Game Archive e não preserva, por si só, o conjunto de materiais primários.")

    cap = doc.add_paragraph("Tabela 3 - Limitações do arquivo de referência, respostas do Press Q e riscos residuais", style="Table Caption")
    table3 = doc.add_table(rows=1, cols=3)
    headers3 = ["Limitação", "Resposta no Press Q", "Risco residual"]
    for idx, text in enumerate(headers3):
        format_cell(table3.rows[0].cells[idx], text, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table3.rows[0])
    rows3 = [
        ("Coleção curada, não repositório integral de fontes primárias", "Declara o corpus como índice vivo e preserva URLs, créditos e notas de evidência.", "Links podem desaparecer; materiais primários ainda exigem preservação institucional e análise de direitos."),
        ("Cobertura incompleta e ausência de denominador", "Restringe percentuais ao corpus e adiciona situação da pesquisa e origem da descoberta.", "O conjunto-semente continua seletivo e não sustenta prevalência ou representatividade da indústria."),
        ("Posts e categorias sobrepostos", "Separa personagens, sistemas e leituras em tabelas com identificadores próprios.", "Relações entre unidades e versões ainda precisam de identificadores persistentes e modelo relacional."),
        ("Evidências de naturezas diferentes", "Registra fonte, confiança, confirmação, limitações e contraevidência.", "Confiança depende de julgamento curatorial; faltam protocolo de codificação e concordância entre avaliadores."),
        ("Interseccionalidade pouco operacionalizada publicamente", "Inclui indicador e detalhes de raça, classe, deficiência, religião e outros eixos quando sustentados.", "O campo atual contém valores inconsistentes; categorias amplas podem reproduzir apagamentos."),
        ("Mudanças de localização, edição e plataforma", "Inclui idioma, plataforma ou versão e data de revisão por registro.", "A versão estava ausente em 46,7% dos personagens; atualizações exigem manutenção contínua."),
        ("Navegação baseada em categorias de site", "Disponibiliza CSV estruturado, análises visuais e consulta por linguagem natural.", "CSV limita colaboração e versionamento; respostas de IA ainda podem simplificar ou errar."),
        ("Trabalho curatorial intensivo e parcialmente não remunerado", "Oferece formulário de contribuição e área administrativa para revisão.", "Participação aumenta a carga de moderação; governança, crédito, consentimento e sustentabilidade financeira permanecem em aberto."),
    ]
    for index, row_data in enumerate(rows3):
        row = table3.add_row()
        for idx, text in enumerate(row_data):
            format_cell(row.cells[idx], text, align=WD_ALIGN_PARAGRAPH.LEFT)
            if index % 2 == 1:
                set_cell_shading(row.cells[idx], PALE_ALT)
    set_table_geometry(table3, [2400, 3100, 3572])
    set_table_borders(table3)
    doc.add_paragraph("Fonte: elaboração própria com base na documentação pública do LGBTQ Video Game Archive e na auditoria do Press Q.", style="Table Source")

    add_heading("5 Discussão", 1)
    add_heading("5.1 Da categoria navegacional à unidade de afirmação", 2)
    add_body("O resultado central é que a principal contribuição do Press Q não está em aumentar a quantidade de categorias, mas em declarar o tipo de afirmação que cada registro sustenta. O LGBTQ Video Game Archive foi desenhado como recurso de descoberta e síntese; suas categorias sobrepostas facilitam percursos temáticos. Quando esse material é reutilizado para visualização quantitativa ou IA, porém, a unidade precisa ser explicitada. A separação entre personagem, sistema e leitura reduz erros de contagem e permite que uma mesma obra apareça em diferentes camadas sem que seus registros sejam tratados como equivalentes.")
    add_body("Essa tipagem também oferece uma solução parcial para o conflito entre canon e recepção. Em vez de excluir interpretações para proteger a precisão ou assimilá-las a identidades confirmadas para preservar sua importância cultural, o modelo mantém uma unidade própria para leituras. A estratégia dialoga com Shaw e Persaud (2020): recepção queer continua pertencendo à história do jogo, mas sua força documental é descrita em termos adequados ao objeto observado.")

    add_heading("5.2 Calculabilidade sem falsa universalidade", 2)
    add_body("Estruturar um corpus torna a contagem mais fácil, mas também aumenta o risco de falsa universalidade. Os percentuais do Press Q são tecnicamente reproduzíveis dentro da captura analisada; ainda assim, o corpus não é amostra probabilística. A concentração de quase metade dos personagens em jogos de 2014 e 2015 demonstra que padrões internos refletem rotas de descoberta e escolhas do conjunto-semente. Por isso, toda visualização deve mostrar denominador, período, unidade e aviso de cobertura.")
    add_body("A distinção de Drucker (2011) entre dados supostamente dados e informação construída é útil aqui. O Press Q não transforma interpretações em fatos neutros ao colocá-las em CSV; ele apenas torna decisões curatoriais mais inspecionáveis. Campos desconhecidos, situações de revisão e confiança precisam aparecer como resultados da análise, não como sujeira a ser eliminada. O achado de valores inconsistentes no indicador de interseccionalidade confirma que a própria estrutura deve ser auditada.")

    add_heading("5.3 Interseccionalidade como requisito de evidência", 2)
    add_body("A expansão para eixos interseccionais responde a uma limitação reconhecida pelo arquivo de referência, mas cria novas responsabilidades. Um esquema pode aumentar visibilidade e simultaneamente consolidar categorias amplas demais. Para evitar isso, o Press Q deveria separar eixos controlados de termos específicos usados pela fonte, permitir múltiplos valores, registrar a justificativa de cada atribuição e proibir inferências por aparência. Revisões com pessoas e comunidades afetadas seriam mais adequadas do que depender exclusivamente de taxonomias internas.")
    add_body("Também é recomendável decompor campos que hoje misturam dimensões. Gênero geral e dimensão trans documentada, por exemplo, não devem formar uma lista de opções mutuamente excludentes. A remodelagem permitiria reconhecer que homens trans são homens e mulheres trans são mulheres sem perder a informação específica relevante para a análise de representação trans.")

    add_heading("5.4 IA como interface, não como autoridade", 2)
    add_body("Quiu amplia acesso ao permitir perguntas em linguagem natural e pode ajudar usuários a localizar evidências e lacunas. Sua arquitetura atual implementa regras importantes: uso exclusivo dos contextos fornecidos, resposta no idioma da pergunta, separação das unidades e obrigação de preservar contraevidência. Essas medidas aproximam o protótipo de princípios de transparência e supervisão humana do NIST (TABASSI, 2023).")
    add_body("Entretanto, instruções não equivalem a verificação. O modelo pode falhar ao contar, omitir uma ressalva ou produzir uma explicação persuasiva para dados inconsistentes. Antes de uso público, recomenda-se um conjunto de testes com perguntas de contagem, ausência, ambiguidade, localização, interseccionalidade e tentativas de induzir inferências. Respostas deveriam ser comparadas a resultados determinísticos e avaliadas por especialistas, usuários e membros das comunidades representadas. Para publicação ou ensino, a fonte primária deve permanecer acessível a partir da resposta.")

    add_heading("5.5 Infraestrutura, crédito e sustentabilidade", 2)
    add_body("A passagem de WordPress para dados estruturados não elimina trabalho; apenas o redistribui. Validação, normalização, revisão, moderação de contribuições, atualização de links e avaliação de IA tornam-se novas tarefas. Um projeto sustentável precisa explicitar papéis, formas de crédito, critérios de aceite, política de correção e recursos institucionais. A contribuição comunitária não deve funcionar como reserva invisível de trabalho gratuito.")
    add_body("Também é necessário reconhecer a dependência intelectual do LGBTQ Video Game Archive. O Press Q deve manter proveniência por registro, citar páginas específicas e devolver correções ou ampliações quando possível. A relação mais produtiva é de interoperabilidade e reciprocidade: o arquivo de referência preserva uma tradição de síntese histórica e leitura contextual, enquanto o Press Q experimenta uma camada analítica estruturada.")

    add_heading("6 Limitações do estudo e agenda de pesquisa", 1)
    add_body("Este estudo possui seis limitações principais. Primeiro, a análise documental concentrou-se nas páginas públicas centrais e em duas publicações associadas; não auditou sistematicamente mais de 1.200 itens da lista-mestra nem entrevistou a equipe do arquivo. Segundo, o corpus do Press Q é pequeno, curado e parcialmente derivado de rotas de descoberta já existentes, o que impede generalizações sobre a indústria ou sobre o próprio LGBTQ Video Game Archive. Terceiro, frequências medem valores registrados, não validade semântica nem qualidade da representação.")
    add_body("Quarto, não houve dupla codificação independente, cálculo de concordância ou validação externa das 94 entradas. Quinto, não foi realizado estudo de usabilidade, acessibilidade ou impacto com pesquisadores, jogadores ou comunidades LGBTQ+. Sexto, a avaliação de Quiu limitou-se à inspeção de código e regras de prompt; não foram medidos fidelidade, alucinação, robustez multilíngue ou vieses. Essas limitações impedem afirmar que o Press Q resolveu os problemas identificados. O protótipo demonstra um desenho possível e torna visível uma agenda de validação.")
    add_body("Como próximos passos, recomenda-se: publicar um protocolo de codificação versionado; migrar para banco de dados com histórico de alterações e identificadores persistentes; realizar amostragem multilíngue e multirregional; implementar revisão independente e adjudicação; normalizar campos interseccionais; publicar dados e documentação com DOI quando direitos permitirem; criar testes determinísticos para a assistente; e conduzir pesquisa participativa sobre terminologia, governança, crédito e dano. Esses passos aproximariam o projeto dos princípios FAIR sem tratar abertura como sinônimo de extração irrestrita.")

    add_heading("7 Considerações finais", 1)
    add_body("O LGBTQ Video Game Archive tornou pesquisável uma história que permaneceria dispersa e vulnerável. Suas limitações mais importantes são amplamente reconhecidas pelo próprio projeto: cobertura incompleta, ausência de denominador, categorias debatidas, fontes heterogêneas, restrições de interseccionalidade e dependência de trabalho curatorial. Analisá-las não significa corrigir um arquivo supostamente falho, mas compreender as condições sob as quais seu conhecimento pode ser reutilizado.")
    add_body("O Press Q oferece uma resposta parcial ao separar identidades de personagens, possibilidades lúdicas e leituras queer. Proveniência, confiança, versão e situação de pesquisa tornam as qualificações mais computáveis; análises visuais e linguagem natural ampliam o acesso. A auditoria mostrou, entretanto, que o próprio protótipo reproduz seletividade, incompletude e inconsistência e acrescenta riscos de automação.")
    add_body("A conclusão é que nenhum esquema encerra a interpretação. A contribuição científica do modelo está em fazer com que diferenças entre tipos de evidência, lacunas e decisões curatoriais sobrevivam à contagem e à consulta por IA. Para arquivos queer, rigor não significa eliminar ambiguidade, mas documentá-la de forma rastreável, contestável e responsável.")

    add_heading("Declarações para submissão", 1)
    declarations = [
        ("Financiamento", "[Informar agência, processo e beneficiário, ou declarar ausência de financiamento específico.]"),
        ("Conflito de interesses", "[Informar conflitos ou declarar que não existem conflitos de interesses.]"),
        ("Contribuições de autoria", "[Preencher conforme a taxonomia CRediT e a política do periódico.]"),
        ("Disponibilidade de dados e código", "[Inserir URL pública e persistente ou DOI do snapshot analisado antes da submissão. Indicar licenças e eventuais restrições.]"),
        ("Ética", "A análise apresentada utiliza documentos, software e registros de obras culturais, sem participantes recrutados. Confirmar a classificação institucional aplicável antes da submissão."),
        ("Uso de IA generativa", "Declaração sugerida: OpenAI Codex foi utilizado como apoio inicial à pesquisa documental, redação e formatação. As fontes, análises, interpretações e a versão final devem ser verificadas e assumidas integralmente pelos autores humanos, conforme a política do periódico."),
    ]
    for label, value in declarations:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(4)
        lead = p.add_run(label + ": ")
        set_run_font(lead, size=10.5, bold=True, color=ACCENT_DARK)
        body = p.add_run(value)
        set_run_font(body, size=10.5, color=INK)

    add_heading("Referências", 1)

    references = [
        ("CRENSHAW, Kimberlé. Demarginalizing the intersection of race and sex: a Black feminist critique of antidiscrimination doctrine, feminist theory and antiracist politics. University of Chicago Legal Forum, v. 1989, n. 1, p. 139-167, 1989. ", "Texto institucional", "https://chicagounbound.uchicago.edu/uclf/vol1989/iss1/8/"),
        ("D'IGNAZIO, Catherine; KLEIN, Lauren F. Data Feminism. Cambridge: MIT Press, 2020. ", "DOI", "https://doi.org/10.7551/mitpress/11805.001.0001"),
        ("DRUCKER, Johanna. Humanities approaches to graphical display. Digital Humanities Quarterly, v. 5, n. 1, 2011. ", "Texto integral", "https://www.digitalhumanities.org/dhq/vol/5/1/000091/000091.html"),
        ("LGBTQ VIDEO GAME ARCHIVE. About (please read first!). Atualização metodológica de 4 jun. 2025. Acesso em: 22 ago. 2026. ", "Página do arquivo", "https://lgbtqgamearchive.com/about/about-archive/"),
        ("LGBTQ VIDEO GAME ARCHIVE. Category descriptions. [s.d.]. Acesso em: 22 ago. 2026. ", "Descrição das categorias", "https://lgbtqgamearchive.com/resources/category-descriptions/"),
        ("LGBTQ VIDEO GAME ARCHIVE. Characters. [s.d.]. Acesso em: 22 ago. 2026. ", "Metodologia de personagens", "https://lgbtqgamearchive.com/resources/category-descriptions/characters/"),
        ("RUBERG, Bonnie. Video Games Have Always Been Queer. New York: New York University Press, 2019. ", "Página da editora", "https://nyupress.org/9781479831036/video-games-have-always-been-queer/"),
        ("RUBERG, Bonnie; SHAW, Adrienne (org.). Queer Game Studies. Minneapolis: University of Minnesota Press, 2017. ", "Página da editora", "https://www.upress.umn.edu/9781517900373/queer-game-studies/"),
        ("SHAW, Adrienne. Gaming at the Edge: Sexuality and Gender at the Margins of Gamer Culture. Minneapolis: University of Minnesota Press, 2014. ", "Página da editora", "https://www.upress.umn.edu/9780816693160/gaming-at-the-edge/"),
        ("SHAW, Adrienne; LAUTERIA, Evan W.; YANG, Hannah; PERSAUD, Christopher J.; COLE, Alayna M. Counting queerness in games: trends in LGBTQ digital game representation, 1985-2005. International Journal of Communication, v. 13, p. 1544-1569, 2019. ", "Texto integral", "https://ijoc.org/index.php/ijoc/article/view/9754/2611"),
        ("SHAW, Adrienne; PERSAUD, Christopher J. Beyond texts: using queer readings to document LGBTQ game content. First Monday, v. 25, n. 8, 2020. ", "DOI", "https://doi.org/10.5210/fm.v25i8.10439"),
        ("TABASSI, Elham. Artificial Intelligence Risk Management Framework (AI RMF 1.0). Gaithersburg: National Institute of Standards and Technology, 2023. NIST AI 100-1. ", "DOI", "https://doi.org/10.6028/NIST.AI.100-1"),
        ("WILKINSON, Mark D. et al. The FAIR Guiding Principles for scientific data management and stewardship. Scientific Data, v. 3, art. 160018, 2016. ", "DOI", "https://doi.org/10.1038/sdata.2016.18"),
    ]
    for text, label, url in references:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(-0.75)
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        set_run_font(r, size=10, color=INK)
        add_hyperlink(p, label, url)

    # Keep captions with objects and prevent widows/orphans where possible.
    for paragraph in doc.paragraphs:
        p_pr = paragraph._p.get_or_add_pPr()
        widow = p_pr.find(qn("w:widowControl"))
        if widow is None:
            widow = OxmlElement("w:widowControl")
            widow.set(qn("w:val"), "true")
            p_pr.append(widow)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
